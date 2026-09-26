"""Exporting the plan: the spreadsheet a coach shares must say what the
planning page says, top down from season to the week's sessions."""
import io
import os
import unittest
from datetime import date, timedelta

os.environ.setdefault("APP_PASSWORD", "test-password")
os.environ.setdefault("SECRET_KEY", "test-secret-key")
os.environ["AI_OPERATION_WORKER_ENABLED"] = "false"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from backend.database import SessionLocal, engine  # noqa: E402
from backend import models  # noqa: E402
from backend.main import app  # noqa: E402
from backend.services.plan_export import build_workbook, phase_fill  # noqa: E402
from backend.tests import reset_database  # noqa: E402

START = date(2026, 9, 7)   # a Monday


class PlanExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        models.Base.metadata.create_all(bind=engine)

    def setUp(self):
        with SessionLocal() as db:
            meet = models.Meet(name="County Champs", date=START + timedelta(weeks=5, days=5))
            db.add(meet)
            db.flush()
            macro = models.TrainingMacro(name="Autumn", squad="Silver 1", sequence_index=1,
                                         date_from=START, date_to=START + timedelta(weeks=6, days=-1),
                                         narrative="Aerobic base into County Champs.", primary_meet_id=meet.id,
                                         group_definitions={"G1": {"description": "Fast lane", "swimmer_ids": [1, 2]}})
            db.add(macro)
            db.flush()
            base = models.SeasonBlock(macro_id=macro.id, sequence_index=1, name="Base", phase_type="base",
                                      date_from=START, date_to=START + timedelta(weeks=4, days=-1),
                                      emphasis={"aerobic": "high"}, group_intents={"G1": "build to 5km"})
            taper = models.SeasonBlock(macro_id=macro.id, sequence_index=2, name="Taper", phase_type="taper",
                                       date_from=START + timedelta(weeks=4), date_to=START + timedelta(weeks=6, days=-1))
            db.add_all([base, taper])
            db.flush()
            db.add(models.Microcycle(
                macro_id=macro.id, block_id=base.id, sequence_index=1, week_start=START,
                week_end=START + timedelta(days=6), label="Week 1 - Base loading", status="confirmed",
                progression_note="Up 10% on last week.",
                sessions=[{"day": "Monday", "date": "2026-09-07", "slot_label": "Mon 06:00", "duration_mins": 90,
                           "session_type": "aerobic", "energy_focus": "aerobic", "key_emphasis": "Long aerobic",
                           "groups": {"G1": {"emphasis": "4x800", "volume_modifier": "full"}}}],
            ))
            for i, value in enumerate([55, 60, 65, 50, 40, 30]):
                db.add(models.SeasonLoadPoint(macro_id=macro.id, week_start=START + timedelta(weeks=i), overall=value,
                                              note="Rest week" if value == 50 else None))
            db.commit()
            self.macro_id = macro.id

    def tearDown(self):
        reset_database()

    def _book(self, **kwargs):
        with SessionLocal() as db:
            content, name = build_workbook(db, **kwargs)
        return load_workbook(io.BytesIO(content)), name

    def test_four_sheets_from_season_down_to_the_week(self):
        book, name = self._book(macro_id=self.macro_id)
        self.assertEqual(book.sheetnames, ["Season calendar", "Macrocycles", "Mesocycles", "Micro layout"])
        self.assertTrue(name.startswith("Autumn-plan-") and name.endswith(".xlsx"))

    def test_the_calendar_has_a_row_per_week_with_its_place_in_the_plan(self):
        ws = self._book(macro_id=self.macro_id)[0]["Season calendar"]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5, max_row=10)]
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[0][2], "1.1.1")
        self.assertEqual(rows[0][3], "Autumn")
        self.assertEqual((rows[0][4], rows[0][5], rows[0][7]), ("Base", "base", 55))
        self.assertEqual(rows[3][8], "Rest week")
        self.assertEqual(rows[4][5], "taper")
        self.assertIn("County Champs", rows[5][9])
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn("D5:D10", merged, "The macro spans its six weeks.")
        self.assertIn("E5:E8", merged, "Base spans its four weeks.")
        self.assertEqual(len(ws._charts), 1, "The planned load is drawn as a line.")
        self.assertIsNotNone(ws["F5"].fill.fgColor.rgb)

    def test_macros_mesos_and_the_micro_layout_carry_their_detail(self):
        book = self._book(macro_id=self.macro_id)[0]
        macro_row = [c.value for c in book["Macrocycles"][5]]
        self.assertEqual(macro_row[1], "Autumn")
        self.assertEqual(macro_row[6], "County Champs")
        self.assertIn("G1: Fast lane (2 swimmers)", macro_row[8])
        meso_rows = [[c.value for c in row] for row in book["Mesocycles"].iter_rows(min_row=5, max_row=6)]
        self.assertEqual([r[0] for r in meso_rows], ["1.1", "1.2"])
        self.assertEqual(meso_rows[0][7], "aerobic: high")
        self.assertEqual(meso_rows[0][8], "G1: build to 5km")
        micro = book["Micro layout"]
        self.assertEqual(micro["C5"].value, "Week 1 - Base loading (confirmed)")
        self.assertIn("Progression: Up 10% on last week.", micro["G5"].value)
        session = [c.value for c in micro[6]]
        self.assertEqual(session[2], "Monday 09-07")
        self.assertEqual(session[3], "Mon 06:00 (90 min)")
        self.assertEqual(session[7], "G1: 4x800 (full)")

    def test_an_empty_plan_still_exports(self):
        reset_database()
        book = self._book()[0]
        self.assertIn("No macrocycles planned yet.", book["Season calendar"]["A2"].value)
        self.assertEqual(book["Micro layout"]["A5"].value, "No weekly plans written yet.")

    def test_phases_are_coloured_by_name(self):
        self.assertIsNotNone(phase_fill("Taper"))
        self.assertIsNotNone(phase_fill("aerobic base"))
        self.assertIsNone(phase_fill(None))

    def test_the_download_is_a_named_spreadsheet(self):
        with TestClient(app) as http:
            token = http.post("/auth/login", json={"password": "test-password"}).json()["token"]
            response = http.get(f"/season/export?macro_id={self.macro_id}",
                                headers={"Authorization": f"Bearer {token}"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.headers["content-type"])
        self.assertIn('filename="Autumn-plan-', response.headers["content-disposition"])


if __name__ == "__main__":
    unittest.main()
