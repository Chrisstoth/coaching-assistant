"""The season plan as a spreadsheet, to share with other coaches.

Four sheets, top down:

    Season calendar  one row per week: macro, meso, micro, phase, planned load,
                     meets and sessions - the layout coaches already plan in -
                     with the macros and mesos spanning their weeks and the
                     load drawn as a line underneath
    Macrocycles      each macro's dates, target meet, intent, groups, pathways
    Mesocycles       each block's phase, dates, emphasis and group intents
    Micro layout     each written week: where it sits, how it progresses, and
                     the session-by-session layout

Built from the same timeline the planning page draws, so the export always
matches what is on screen.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session as DBSession

from backend import models

# Light fills so a printout stays readable. Matched loosely on the phase name.
PHASE_FILLS = [
    ("taper", "E9D5FF"),
    ("peak", "FECACA"),
    ("compet", "BBF7D0"),
    ("race", "BBF7D0"),
    ("recover", "E5E7EB"),
    ("transition", "E5E7EB"),
    ("rest", "E5E7EB"),
    ("build", "FED7AA"),
    ("specific", "FED7AA"),
    ("base", "BFDBFE"),
    ("aerobic", "BFDBFE"),
    ("general", "BFDBFE"),
    ("prep", "BFDBFE"),
]
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WEEK_FILL = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)


def phase_fill(phase: Optional[str]) -> Optional[PatternFill]:
    lowered = (phase or "").lower()
    for word, colour in PHASE_FILLS:
        if word in lowered:
            return PatternFill("solid", fgColor=colour)
    return None


def _text(value) -> str:
    """Stored JSON (emphasis, intents) as plain reading text."""
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_text(v)}" for k, v in value.items() if _text(v))
    if isinstance(value, list):
        return ", ".join(_text(v) for v in value if _text(v))
    return str(value)


def _day(value) -> Optional[date]:
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def _header(ws, row: int, titles: list, widths: list) -> None:
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, HEADER_FONT, CENTRE, BOX
        ws.column_dimensions[get_column_letter(col)].width = width


def _title(ws, text: str, subtitle: str) -> None:
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=subtitle).font = Font(italic=True, color="6B7280")


def _merge_runs(ws, column: int, first_row: int, values: list) -> None:
    """Merge consecutive equal, non-empty cells down a column: a macro or meso
    then spans its weeks the way it does on a planning wall chart."""
    start = 0
    for i in range(1, len(values) + 1):
        if i == len(values) or values[i] != values[start]:
            if values[start] and i - start > 1:
                ws.merge_cells(start_row=first_row + start, start_column=column,
                               end_row=first_row + i - 1, end_column=column)
            start = i


def _calendar_sheet(ws, timeline: dict, heading: str) -> None:
    weeks = timeline["weeks"]
    _title(ws, heading, f"{timeline['date_from']} to {timeline['date_to']} - exported {date.today():%d %b %Y}")
    titles = ["Week of", "Wk", "Code", "Macrocycle", "Mesocycle", "Phase", "Micro",
              "Planned load (0-100)", "Load note", "Meets", "Sessions"]
    _header(ws, 4, titles, [11, 9, 8, 22, 22, 13, 7, 11, 28, 34, 12])
    first = 5
    for i, week in enumerate(weeks):
        r = first + i
        load = week.get("load") or {}
        sessions = week.get("sessions") or {}
        session_text = f"{sessions.get('planned', 0)}" + (
            f" ({sessions['cancelled']} cancelled)" if sessions.get("cancelled") else "")
        values = [
            _day(week["week_start"]), week.get("iso_week", "").split("-")[-1], week.get("cycle_code") or "",
            week.get("macro_name") or "", week.get("block_name") or "", week.get("phase_type") or "",
            week.get("micro_index") or "", load.get("overall"), load.get("note") or "",
            "; ".join(f"{_day(m['date']):%d %b} {m['name']}" for m in week.get("meets") or []),
            session_text if (sessions.get("planned") or sessions.get("cancelled")) else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border, cell.alignment = BOX, WRAP
        ws.cell(row=r, column=1).number_format = "dd mmm yy"
        fill = phase_fill(week.get("phase_type"))
        if fill:
            for col in (5, 6, 7):
                ws.cell(row=r, column=col).fill = fill
        if week.get("is_current"):
            for col in range(1, 4):
                ws.cell(row=r, column=col).font = Font(bold=True)
        if week.get("meets"):
            ws.cell(row=r, column=10).font = Font(bold=True, color="B91C1C")

    _merge_runs(ws, 4, first, [w.get("macro_name") or "" for w in weeks])
    _merge_runs(ws, 5, first, [f"{w.get('macro_id')}:{w.get('block_id')}" if w.get("block_id") else "" for w in weeks])
    for col in (4, 5):
        for i in range(len(weeks)):
            ws.cell(row=first + i, column=col).alignment = CENTRE
    ws.freeze_panes = ws.cell(row=first, column=2)

    if any((w.get("load") or {}).get("overall") is not None for w in weeks):
        chart = LineChart()
        chart.title = "Planned weekly load"
        chart.y_axis.title = "Load (0-100)"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 100
        chart.height, chart.width = 7.5, max(16, len(weeks) * 0.6)
        chart.legend = None
        data = Reference(ws, min_col=8, min_row=4, max_row=first + len(weeks) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=first + len(weeks) - 1))
        chart.x_axis.number_format = "dd mmm"
        ws.add_chart(chart, f"A{first + len(weeks) + 2}")


def _macro_sheet(ws, db: DBSession, macros: list) -> None:
    _title(ws, "Macrocycles", "The big blocks of the season, each built towards a target meet.")
    _header(ws, 4, ["#", "Macrocycle", "Squad", "From", "To", "Weeks", "Target meet", "Intent", "Groups", "Pathways"],
            [5, 24, 14, 11, 11, 7, 26, 50, 40, 50])
    for i, macro in enumerate(macros):
        r = 5 + i
        groups = "; ".join(
            f"{name}: {defn.get('description', '')}".strip(": ") + (
                f" ({len(defn.get('swimmer_ids') or [])} swimmers)" if defn.get("swimmer_ids") else "")
            for name, defn in (macro.group_definitions or {}).items() if isinstance(defn, dict))
        pathways = []
        for p in db.query(models.PlanningPathway).filter(
                models.PlanningPathway.macro_id == macro.id, models.PlanningPathway.active.is_(True)).all():
            members = len([m for m in p.memberships if m.active])
            target = p.primary_meet.name if p.primary_meet else "no target meet"
            fallback = f", else {p.fallback_meet.name}" if p.fallback_meet else ""
            pathways.append(f"{p.name} -> {target}{fallback} ({members} swimmers)")
        values = [macro.sequence_index or i + 1, macro.name, macro.squad or "", macro.date_from, macro.date_to,
                  round(((macro.date_to - macro.date_from).days + 1) / 7),
                  macro.primary_meet.name if macro.primary_meet else "", macro.narrative or "", groups,
                  "; ".join(pathways)]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border, cell.alignment = BOX, WRAP
        for col in (4, 5):
            ws.cell(row=r, column=col).number_format = "dd mmm yy"
    ws.freeze_panes = "C5"


def _meso_sheet(ws, blocks: list) -> None:
    _title(ws, "Mesocycles", "The blocks inside each macrocycle, with their phase and emphasis.")
    _header(ws, 4, ["Code", "Macrocycle", "Mesocycle", "Phase", "From", "To", "Weeks", "Emphasis", "Group intents", "Notes"],
            [8, 22, 24, 13, 11, 11, 7, 40, 50, 40])
    for i, block in enumerate(blocks):
        r = 5 + i
        macro = block.macro
        code = f"{macro.sequence_index}.{block.sequence_index}" if macro else str(block.sequence_index or "")
        values = [code, macro.name if macro else "", block.name, block.phase_type or "", block.date_from,
                  block.date_to, round(((block.date_to - block.date_from).days + 1) / 7),
                  _text(block.emphasis), _text(block.group_intents), block.notes or ""]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border, cell.alignment = BOX, WRAP
        for col in (5, 6):
            ws.cell(row=r, column=col).number_format = "dd mmm yy"
        fill = phase_fill(block.phase_type)
        if fill:
            for col in (3, 4):
                ws.cell(row=r, column=col).fill = fill
    ws.freeze_panes = "D5"


def _micro_sheet(ws, micros: list, code_by_week: dict) -> None:
    _title(ws, "Micro layout", "Each written week, then its sessions in order.")
    _header(ws, 4, ["Week of", "Code", "Day", "Slot", "Type", "Focus", "Emphasis", "Groups", "Notes"],
            [11, 8, 11, 14, 13, 13, 40, 50, 30])
    r = 5
    if not micros:
        ws.cell(row=r, column=1, value="No weekly plans written yet.")
        return
    for micro in micros:
        week_cells = [micro.week_start, code_by_week.get((micro.block_id, micro.week_start), ""),
                      f"{micro.label} ({micro.status})"]
        for col, value in enumerate(week_cells, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.fill, cell.border = Font(bold=True), WEEK_FILL, BOX
        ws.cell(row=r, column=1).number_format = "dd mmm yy"
        notes = " | ".join(f"{label}: {text}" for label, text in (
            ("Where it sits", micro.meso_position_note), ("Progression", micro.progression_note),
            ("Recovery", micro.recovery_placement), ("Next week", micro.next_week_direction)) if text)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        note_cell = ws.cell(row=r, column=7, value=notes)
        note_cell.alignment, note_cell.fill = WRAP, WEEK_FILL
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
        r += 1
        for session in micro.sessions or []:
            if not isinstance(session, dict):
                continue
            groups = "; ".join(
                f"{name}: {g.get('emphasis', '')}" + (f" ({g['volume_modifier']})" if g.get("volume_modifier") else "")
                for name, g in (session.get("groups") or {}).items() if isinstance(g, dict))
            slot = session.get("slot_label") or ""
            if session.get("duration_mins"):
                slot = f"{slot} ({session['duration_mins']} min)".strip()
            values = ["", "", f"{session.get('day', '')} {session.get('date', '')[5:] if session.get('date') else ''}".strip(),
                      slot, session.get("session_type") or "", session.get("energy_focus") or "",
                      session.get("key_emphasis") or "", groups, session.get("pool_note") or ""]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=value)
                cell.border, cell.alignment = BOX, WRAP
            r += 1
    ws.freeze_panes = "C5"


def build_workbook(db: DBSession, macro_id: Optional[int] = None, squad: Optional[str] = None) -> tuple:
    """(xlsx bytes, file name)."""
    from backend.routers.season import get_timeline

    timeline = get_timeline(macro_id=macro_id, squad=squad, pathway_id=None, db=db)
    macro_q = db.query(models.TrainingMacro)
    if macro_id:
        macro_q = macro_q.filter(models.TrainingMacro.id == macro_id)
    if squad:
        macro_q = macro_q.filter(models.TrainingMacro.squad == squad)
    macros = macro_q.order_by(models.TrainingMacro.date_from).all()
    macro_ids = [m.id for m in macros]
    blocks = db.query(models.SeasonBlock).filter(models.SeasonBlock.macro_id.in_(macro_ids)).order_by(
        models.SeasonBlock.date_from).all() if macro_ids else []
    block_ids = [b.id for b in blocks]
    micros = db.query(models.Microcycle).filter(models.Microcycle.block_id.in_(block_ids)).order_by(
        models.Microcycle.week_start).all() if block_ids else []
    code_by_week = {(w["block_id"], _day(w["week_start"])): w["cycle_code"] for w in timeline["weeks"] if w.get("block_id")}

    if macro_id and macros:
        heading = f"{macros[0].name} - training plan"
        name = macros[0].name
    else:
        heading = f"{squad + ' - ' if squad else ''}Season plan"
        name = squad or "season"

    wb = Workbook()
    calendar = wb.active
    calendar.title = "Season calendar"
    if timeline["weeks"]:
        _calendar_sheet(calendar, timeline, heading)
    else:
        _title(calendar, heading, "No macrocycles planned yet.")
    _macro_sheet(wb.create_sheet("Macrocycles"), db, macros)
    _meso_sheet(wb.create_sheet("Mesocycles"), blocks)
    _micro_sheet(wb.create_sheet("Micro layout"), micros, code_by_week)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "4:4"

    buffer = io.BytesIO()
    wb.save(buffer)
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "" for ch in name).strip().replace(" ", "-") or "plan"
    return buffer.getvalue(), f"{safe}-plan-{date.today():%Y-%m-%d}.xlsx"
